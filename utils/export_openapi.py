#!/usr/bin/env python3
"""
OpenAPI Schema Snapshot & Documentation Exporter (CLI Tool).

This tool extracts the OpenAPI 3.x schema definition from a FastAPI application
(via direct app import, live HTTP endpoint, or static JSON file) and exports
a comprehensive API Catalog into CSV and formatted Excel (.xlsx) formats.

Usage Examples:
---------------
1. Direct from FastAPI App (Default):
   python utils/export_openapi.py --from-app

2. From running server endpoint:
   python utils/export_openapi.py --url http://localhost:8000/openapi.json

3. From local openapi.json file:
   python utils/export_openapi.py --file backend/openapi.json

Outputs:
--------
- api_snapshot.csv  : Comma-separated flat API endpoint inventory.
- api_snapshot.xlsx : Formatted multi-sheet Excel report with method badges,
                      auto-fitted columns, and categorical summary statistics.
"""

import os
import sys
import json
import argparse
from datetime import datetime
from pathlib import Path
from typing import Dict, Any, List

# Ensure backend path is discoverable if importing directly
CURRENT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = CURRENT_DIR.parent
BACKEND_DIR = PROJECT_ROOT / "backend"
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))


def load_schema_from_app() -> Dict[str, Any]:
    """Extract OpenAPI schema directly by importing the FastAPI instance."""
    print("[INFO] Importing FastAPI application from 'main.py'...")
    try:
        from main import app
        schema = app.openapi()
        print(f"[SUCCESS] Generated OpenAPI schema for: '{schema.get('info', {}).get('title')}' v{schema.get('info', {}).get('version')}")
        return schema
    except ImportError as e:
        print(f"[ERROR] Could not import FastAPI app: {e}")
        sys.exit(1)


def load_schema_from_url(url: str) -> Dict[str, Any]:
    """Fetch OpenAPI schema JSON from a running HTTP server endpoint."""
    import requests
    print(f"[INFO] Fetching OpenAPI schema from URL: {url}...")
    try:
        response = requests.get(url, timeout=10)
        response.raise_for_status()
        schema = response.json()
        print(f"[SUCCESS] Successfully retrieved schema from {url}")
        return schema
    except Exception as e:
        print(f"[ERROR] Failed to fetch schema from {url}: {e}")
        sys.exit(1)


def load_schema_from_file(file_path: str) -> Dict[str, Any]:
    """Read OpenAPI schema from a local JSON file."""
    print(f"[INFO] Reading schema from file: {file_path}...")
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            schema = json.load(f)
        print(f"[SUCCESS] Loaded schema file from {file_path}")
        return schema
    except Exception as e:
        print(f"[ERROR] Failed to read file {file_path}: {e}")
        sys.exit(1)


def parse_openapi_endpoints(schema: Dict[str, Any]) -> List[Dict[str, Any]]:
    """
    Parses OpenAPI 3.x schema and flattens endpoints into structured dictionary records.
    """
    endpoints = []
    paths = schema.get("paths", {})
    http_methods = ["get", "post", "put", "delete", "patch", "options", "head"]

    for path_url, path_item in paths.items():
        for method_key in http_methods:
            if method_key not in path_item:
                continue

            op = path_item[method_key]
            method_str = method_key.upper()
            tags = ", ".join(op.get("tags", ["Default"]))
            summary = op.get("summary", "")
            description = (op.get("description") or "").strip().replace("\n", " ")
            operation_id = op.get("operationId", "")
            deprecated = "Yes" if op.get("deprecated", False) else "No"

            # Parse parameters (query, path, header, cookie)
            param_list = []
            for p in op.get("parameters", []):
                p_name = p.get("name", "")
                p_in = p.get("in", "")
                p_req = "Required" if p.get("required", False) else "Optional"
                p_type = p.get("schema", {}).get("type", "any")
                param_list.append(f"{p_name} ({p_in}, {p_type}, {p_req})")
            parameters_str = "; ".join(param_list) if param_list else "-"

            # Parse request body
            request_body = op.get("requestBody", {})
            rb_content = request_body.get("content", {})
            rb_types = list(rb_content.keys())
            request_body_str = ", ".join(rb_types) if rb_types else "-"

            # Parse responses
            responses = op.get("responses", {})
            resp_list = []
            for status_code, resp_obj in responses.items():
                resp_desc = resp_obj.get("description", "")
                resp_list.append(f"HTTP {status_code}: {resp_desc}")
            responses_str = " | ".join(resp_list) if resp_list else "-"

            endpoints.append({
                "Tag / Module": tags,
                "HTTP Method": method_str,
                "Endpoint Path": path_url,
                "Summary": summary,
                "Description": description,
                "Parameters": parameters_str,
                "Request Body": request_body_str,
                "Responses": responses_str,
                "Operation ID": operation_id,
                "Deprecated": deprecated
            })

    return endpoints


def export_to_csv(endpoints: List[Dict[str, Any]], output_path: str):
    """Exports parsed endpoints list to a CSV file."""
    import csv
    if not endpoints:
        print("[WARNING] No endpoints found to export.")
        return

    headers = list(endpoints[0].keys())
    with open(output_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        writer.writerows(endpoints)

    print(f"[SUCCESS] CSV Exported: {output_path} ({len(endpoints)} endpoints)")


def export_to_excel(endpoints: List[Dict[str, Any]], schema_info: Dict[str, Any], output_path: str):
    """Exports parsed endpoints to a styled multi-sheet Excel workbook using openpyxl."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    
    # Sheet 1: API Endpoint Catalog
    ws = wb.active
    ws.title = "API Endpoints"
    ws.views.sheetView[0].showGridLines = True

    # Styling Palette
    FONT_FAMILY = "Segoe UI"
    title_font = Font(name=FONT_FAMILY, size=16, bold=True, color="1F4E79")
    subtitle_font = Font(name=FONT_FAMILY, size=10, italic=True, color="595959")
    header_font = Font(name=FONT_FAMILY, size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    
    method_colors = {
        "GET": PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid"),     # Soft Green
        "POST": PatternFill(start_color="CFE2F3", end_color="CFE2F3", fill_type="solid"),    # Soft Blue
        "PUT": PatternFill(start_color="FCE5CD", end_color="FCE5CD", fill_type="solid"),     # Soft Orange
        "DELETE": PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid"),  # Soft Red
        "PATCH": PatternFill(start_color="EAD1DC", end_color="EAD1DC", fill_type="solid"),   # Soft Purple
    }
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # 1. Title Banner
    api_title = schema_info.get("title", "FastAPI Service")
    api_version = schema_info.get("version", "1.0.0")
    ws.merge_cells("A1:J1")
    ws["A1"] = f"{api_title} - API Catalog Snapshot (v{api_version})"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(vertical="center")

    ws.merge_cells("A2:J2")
    ws["A2"] = f"Generated at: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Total Endpoints: {len(endpoints)}"
    ws["A2"].font = subtitle_font
    ws["A2"].alignment = Alignment(vertical="center")

    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[4].height = 26

    # 2. Table Headers
    if endpoints:
        headers = list(endpoints[0].keys())
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=4, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        # 3. Data Rows
        for row_idx, ep in enumerate(endpoints, start=5):
            ws.row_dimensions[row_idx].height = 22
            for col_idx, header in enumerate(headers, start=1):
                val = ep.get(header, "")
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = Font(name=FONT_FAMILY, size=10)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

                # Method badge styling
                if header == "HTTP Method" and val in method_colors:
                    cell.fill = method_colors[val]
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.font = Font(name=FONT_FAMILY, size=10, bold=True)

        # 4. Auto-fit column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row < 4:
                    continue
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 50)

    # Sheet 2: Summary Metrics
    ws_summary = wb.create_sheet(title="Summary Statistics")
    ws_summary.views.sheetView[0].showGridLines = True
    
    ws_summary["A1"] = "Category"
    ws_summary["B1"] = "Count"
    ws_summary["A1"].font = header_font
    ws_summary["A1"].fill = header_fill
    ws_summary["B1"].font = header_font
    ws_summary["B1"].fill = header_fill

    method_counts: Dict[str, int] = {}
    tag_counts: Dict[str, int] = {}
    for ep in endpoints:
        m = ep["HTTP Method"]
        t = ep["Tag / Module"]
        method_counts[m] = method_counts.get(m, 0) + 1
        tag_counts[t] = tag_counts.get(t, 0) + 1

    r = 2
    ws_summary.cell(row=r, column=1, value="Total Endpoints").font = Font(name=FONT_FAMILY, bold=True)
    ws_summary.cell(row=r, column=2, value=len(endpoints)).font = Font(name=FONT_FAMILY, bold=True)
    r += 2

    ws_summary.cell(row=r, column=1, value="Endpoints by Method").font = Font(name=FONT_FAMILY, bold=True, color="1F4E79")
    r += 1
    for m, c in sorted(method_counts.items()):
        ws_summary.cell(row=r, column=1, value=m).font = Font(name=FONT_FAMILY)
        ws_summary.cell(row=r, column=2, value=c).font = Font(name=FONT_FAMILY)
        r += 1

    r += 1
    ws_summary.cell(row=r, column=1, value="Endpoints by Module / Tag").font = Font(name=FONT_FAMILY, bold=True, color="1F4E79")
    r += 1
    for t, c in sorted(tag_counts.items()):
        ws_summary.cell(row=r, column=1, value=t).font = Font(name=FONT_FAMILY)
        ws_summary.cell(row=r, column=2, value=c).font = Font(name=FONT_FAMILY)
        r += 1

    ws_summary.column_dimensions["A"].width = 35
    ws_summary.column_dimensions["B"].width = 15

    wb.save(output_path)
    print(f"[SUCCESS] Excel Report Exported: {output_path}")


def main():
    parser = argparse.ArgumentParser(
        description="Extract OpenAPI specification and export into CSV and formatted Excel (.xlsx) files."
    )
    group = parser.add_mutually_exclusive_group()
    group.add_argument("--from-app", action="store_true", default=True, help="Load directly from FastAPI app instance (default)")
    group.add_argument("--url", type=str, help="Fetch OpenAPI schema from URL (e.g. http://localhost:8000/openapi.json)")
    group.add_argument("--file", type=str, help="Load schema from a local JSON file path")

    parser.add_argument("--output-dir", type=str, default=".", help="Directory to save export files (default: current directory)")
    parser.add_argument("--save-json", action="store_true", default=True, help="Also save openapi.json file snapshot")

    args = parser.parse_args()

    # Determine schema loading source
    if args.url:
        schema = load_schema_from_url(args.url)
    elif args.file:
        schema = load_schema_from_file(args.file)
    else:
        schema = load_schema_from_app()

    schema_info = schema.get("info", {})
    endpoints = parse_openapi_endpoints(schema)
    print(f"[INFO] Parsed {len(endpoints)} API operations across {len(schema.get('paths', {}))} paths.")

    out_dir = Path(args.output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    # 1. Save openapi.json
    if args.save_json:
        json_path = out_dir / "openapi.json"
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(schema, f, indent=2, ensure_ascii=False)
        print(f"[SUCCESS] Saved openapi.json snapshot: {json_path}")

    # 2. Export CSV
    csv_path = out_dir / "api_snapshot.csv"
    export_to_csv(endpoints, str(csv_path))

    # 3. Export Excel
    excel_path = out_dir / "api_snapshot.xlsx"
    export_to_excel(endpoints, schema_info, str(excel_path))

    print("\n" + "=" * 60)
    print("Snapshot Generation Summary:")
    print(f"  • OpenAPI JSON: {out_dir / 'openapi.json'}")
    print(f"  • CSV Catalog:  {csv_path}")
    print(f"  • Excel Report: {excel_path}")
    print("=" * 60)


if __name__ == "__main__":
    main()
